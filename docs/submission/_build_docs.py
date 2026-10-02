from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
LOGO = OUT / "nirma_logo.png"
NAVY = "203451"
BLUE = "315B84"
GOLD = "C49A45"
PALE = "F2F5F8"
PALE_GOLD = "FAF5E8"
GRAY = "596579"
RED = "8E3040"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_keep_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_run_font(run, name="Arial", size=10, color=NAVY, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    set_run_font(run, size=8, color=GRAY)


def setup_doc(short_title):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size, color in (("Heading 1", 18, NAVY), ("Heading 2", 13, BLUE), ("Heading 3", 11, RED)):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    title_style = styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_style.font.size = Pt(27)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string(NAVY)
    title_style.paragraph_format.space_after = Pt(9)
    title_style.paragraph_format.keep_with_next = True

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("NIRMA UNIVERSITY  |  4CS101ME25"), size=8, color=GRAY, bold=True)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run(f"{short_title}   •   Page "), size=8, color=GRAY)
    add_page_number(footer)
    return doc


def cover(doc, title, subtitle, doc_type, include_identity=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(10)
    p.add_run().add_picture(str(LOGO), width=Inches(2.15))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run("NIRMA UNIVERSITY"), size=15, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(17)
    set_run_font(p.add_run("INSTITUTE OF TECHNOLOGY"), size=10, color=GRAY, bold=True)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(17)
    set_run_font(p.add_run(subtitle), size=12, color=BLUE)
    info = doc.add_table(rows=3, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    info.columns[0].width = Inches(1.65)
    info.columns[1].width = Inches(4.35)
    rows = [
        ("Subject", "4CS101ME25 - Big Data Systems"),
        ("Document", doc_type),
        ("Prepared", "2 October 2026"),
    ]
    for row, (key, value) in zip(info.rows, rows):
        for cell in row.cells:
            set_cell_margins(cell, top=120, bottom=120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(row.cells[0], NAVY)
        shade(row.cells[1], PALE)
        set_run_font(row.cells[0].paragraphs[0].add_run(key), size=9, color="FFFFFF", bold=True)
        set_run_font(row.cells[1].paragraphs[0].add_run(value), size=9, color=NAVY, bold=True)
        set_keep_row(row)
    if include_identity:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run("Submitted by: Fenil Chodvadiya (Lead), Sarth Narola, Aayush Savaliya"), size=9.5, color=NAVY, bold=True)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run("Department of Computer Science & Engineering  •  Institute of Technology"), size=9, color=GRAY)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run("Subject: 4CS101ME25 - Big Data Systems"), size=9, color=BLUE, bold=True)
    doc.add_page_break()


def para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), bold=True)
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)
    return p


def number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_together = True
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F0F3F6")
    pPr.append(shd)
    for i, line in enumerate(text.strip("\n").splitlines()):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, "Consolas", 8.2, NAVY)
    return p


def callout(doc, label, text, fill=PALE_GOLD):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, top=145, start=180, bottom=145, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    set_run_font(p.add_run(label + "  "), size=9, color=RED, bold=True)
    set_run_font(p.add_run(text), size=9, color=NAVY)
    set_keep_row(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def table(doc, headers, rows, widths=None, font_size=8.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        if widths:
            c.width = Inches(widths[i])
        shade(c, NAVY)
        set_cell_margins(c)
        set_run_font(c.paragraphs[0].add_run(h), size=font_size, color="FFFFFF", bold=True)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_repeat_table_header(t.rows[0])
    set_keep_row(t.rows[0])
    for ri, row_data in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row_data):
            if widths:
                cells[i].width = Inches(widths[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2 == 1:
                shade(cells[i], "F7F9FB")
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            set_run_font(p.add_run(str(value)), size=font_size, color=NAVY)
        set_keep_row(t.rows[-1])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def build_run_guide():
    doc = setup_doc("Run guide")
    cover(doc, "Run the Project Step by Step", "A cross-platform command guide for the COVID-19 Big Data Analytics project", "System run guide", False)
    heading(doc, "Purpose and verified operating model")
    para(doc, "Use this guide to prepare a new computer, start the Docker services, place the source files in HDFS, submit the Spark job, inspect its result, and open Streamlit. Hadoop and Spark run in Docker; Python on the host is used for data upload and optional local dashboard work.")
    callout(doc, "Important status", "The cluster and HDFS input were verified on 2 October 2026. A fresh Spark run registered executors on both workers but lost an executor during shuffle and did not finish. Do not treat existing output/ files as freshly generated until a new run exits successfully.")
    heading(doc, "1. Install the tools")
    bullet(doc, "Install Git and Docker Desktop (Windows/macOS) or Docker Engine plus Compose v2 (Linux). Start the Docker service before continuing.")
    bullet(doc, "Install Python 3.10 or 3.11 if you will use the host upload scripts or run Streamlit outside Docker. Java and Hadoop do not need host installation.")
    code(doc, "git --version\ndocker --version\ndocker compose version\npy --version")
    heading(doc, "2. Clone the repository")
    heading(doc, "Windows PowerShell", 2)
    code(doc, "git clone https://github.com/Fenil412/covid-big-data-analytics.git\nSet-Location covid-big-data-analytics")
    heading(doc, "Windows Command Prompt", 2)
    code(doc, "git clone https://github.com/Fenil412/covid-big-data-analytics.git\ncd covid-big-data-analytics")
    heading(doc, "Linux or macOS", 2)
    code(doc, "git clone https://github.com/Fenil412/covid-big-data-analytics.git\ncd covid-big-data-analytics")
    heading(doc, "3. Create local configuration")
    para(doc, "Compose requires a local .env file. Copy the example and edit it on your own machine. For Parquet-only use, set MONGODB_ATLAS_URI to an empty value. Add your own MongoDB Atlas URI only if you intend to use MongoDB. Never commit .env or paste credentials into screenshots.")
    heading(doc, "Windows PowerShell", 2)
    code(doc, "Copy-Item .env.example .env\nnotepad .env")
    heading(doc, "Windows Command Prompt", 2)
    code(doc, "copy .env.example .env\nnotepad .env")
    heading(doc, "Linux or macOS", 2)
    code(doc, "cp .env.example .env\nnano .env")
    heading(doc, "4. Optional host virtual environment")
    para(doc, "Use the project environment for the source downloader/uploader and optional local Streamlit. Docker installs dashboard packages from requirements-dashboard.txt when it builds the dashboard image.")
    heading(doc, "Windows PowerShell", 2)
    code(doc, "py -3.11 -m venv .venv\n.\\.venv\\Scripts\\Activate.ps1\npython -m pip install --upgrade pip\npip install -r requirements.txt")
    heading(doc, "Windows Command Prompt", 2)
    code(doc, "py -3.11 -m venv .venv\n.venv\\Scripts\\activate.bat\npython -m pip install --upgrade pip\npip install -r requirements.txt")
    heading(doc, "Linux or macOS", 2)
    code(doc, "python3 -m venv .venv\nsource .venv/bin/activate\npython -m pip install --upgrade pip\npip install -r requirements.txt")
    heading(doc, "5. Build and start the complete Docker stack")
    code(doc, "docker compose config --quiet\ndocker compose build\ndocker compose up -d\ndocker compose ps")
    para(doc, "The Compose project starts the NameNode, ResourceManager, two DataNodes, two NodeManagers, Spark master, two Spark workers, and dashboard. The first image downloads can take several minutes.")
    heading(doc, "6. Verify HDFS and the two worker nodes")
    code(doc, "docker exec hadoop-master hdfs dfsadmin -report\ndocker exec hadoop-master hdfs dfs -ls /covid/input\ndocker exec hadoop-master hdfs fsck /covid/input/epidemiology.csv -files -blocks -locations")
    callout(doc, "Expected check", "The HDFS report should show two live DataNodes. For the verified epidemiology input, fsck showed four healthy blocks at replication factor two, with locations on both workers. The project source file is epidemiology.csv; the sample name covid.csv is not used here.", PALE)
    heading(doc, "7. Download the source data and upload it to HDFS")
    para(doc, "Activate the virtual environment first, then run the helper from the repository root. The helper downloads the Google COVID-19 Open Data inputs to dataset/ and copies them into HDFS.")
    table(doc, ["Platform", "Command"], [
        ("Windows CMD or PowerShell", "upload-data.bat"),
        ("Linux/macOS", "bash scripts/cluster/upload-data.sh"),
    ], [2.05, 4.45])
    code(doc, "docker exec hadoop-master hdfs dfs -ls /covid/input\ndocker exec hadoop-master hdfs fsck /covid/input/epidemiology.csv -files -blocks -locations")
    heading(doc, "8. Submit the distributed analytics job")
    para(doc, "Use standalone mode to submit to the two Spark worker containers. Current Hadoop NodeManager images do not include Python 3 for YARN worker execution.")
    table(doc, ["Platform", "Command"], [
        ("Windows CMD", "run-analysis.bat -Mode cluster"),
        ("Windows PowerShell", "powershell -ExecutionPolicy Bypass -File .\\scripts\\cluster\\run-analysis.ps1 -Mode cluster"),
        ("Linux/macOS", "SPARK_RUN_MODE=cluster bash scripts/cluster/run-analysis.sh"),
    ], [2.05, 4.45], 8)
    code(doc, "docker compose logs --tail=100 spark-master spark-worker1 spark-worker2\ndocker exec hadoop-master hdfs dfs -ls -R /covid/output")
    callout(doc, "Only accept success when", "The Spark submit command exits with code 0, expected Parquet result directories exist in HDFS, and the job date/schema are current. The last verified attempt lost an executor during shuffle; a listed output folder by itself does not prove success.")
    heading(doc, "9. Open the Streamlit dashboard")
    para(doc, "The Compose dashboard is available at http://localhost:8501. It reads processed MongoDB collections when configured and available, then falls back to local output/ Parquet files. It is a visualization and exploration layer, not the Spark processing engine.")
    code(doc, "# Optional: run locally after activating .venv\npython -m streamlit run dashboard/app.py")
    heading(doc, "10. Stop services and preserve or reset HDFS")
    code(doc, "# Stop containers while retaining named HDFS volumes\ndocker compose down")
    callout(doc, "Destructive reset", "docker compose down --volumes removes the HDFS data and NameNode metadata. Run it only when you deliberately want an empty cluster.", "F9EDEF")
    heading(doc, "Quick troubleshooting")
    table(doc, ["Symptom", "First checks"], [
        ("Compose says .env is missing", "Copy .env.example to .env and keep MONGODB_ATLAS_URI empty for Parquet-only mode."),
        ("Fewer than two live DataNodes", "Check docker compose ps and docker compose logs datanode1 datanode2."),
        ("HDFS input is empty", "Run upload-data.bat or scripts/cluster/upload-data.sh with the venv active."),
        ("Spark exits 137 or executor disappears", "Inspect Spark logs and Docker memory; do not call stale Parquet a new result."),
        ("YARN Python worker error", "Use Spark standalone mode, or add Python 3 to the NodeManager images."),
        ("Dashboard has no analytics", "Check output/ files or the configured MongoDB collections and their dates."),
    ], [2.0, 4.5], 8)
    return doc


def build_explainer():
    doc = setup_doc("Professor briefing")
    cover(doc, "Explain the Project to Your Professor", "Architecture, processing flow, demonstration script, and viva preparation", "Project explanation and presentation guide")
    heading(doc, "Project in one minute")
    para(doc, "This project demonstrates a COVID-19 analytics pipeline built around distributed storage and processing. Five CSV inputs are uploaded into HDFS, where the large epidemiology file is divided into blocks and replicated across two DataNodes. A PySpark job reads those HDFS inputs and prepares country, global, date, vaccination, hospitalization, and regional summaries. Streamlit presents processed summaries for exploration; it does not replace Spark or perform the distributed aggregation.")
    callout(doc, "Accurate cluster description", "The Hadoop cluster has one master and two worker containers. They run on one Docker host, so this is a containerized distributed cluster demonstration, not three separate physical computers.", PALE)
    heading(doc, "What each technology contributes")
    table(doc, ["Component", "Role in this project", "What to point out"], [
        ("Docker Compose", "Starts and connects the project services.", "One master, two Hadoop workers, Spark master/two workers, and Streamlit."),
        ("HDFS", "Stores source files as distributed, replicated blocks.", "Use dfsadmin -report and fsck locations as evidence."),
        ("YARN", "Hadoop resource management services are present.", "ResourceManager and NodeManagers run; current job path uses Spark standalone."),
        ("Apache Spark", "Processes the COVID tables and prepares aggregate DataFrames.", "The latest attempt registered an executor on each Spark worker but did not finish output."),
        ("MongoDB / Parquet", "Optional results destinations and dashboard fallback.", "Do not label an old or partial result as current."),
        ("Streamlit", "Shows KPIs, country views, charts, filters, and data quality.", "It consumes results; it is not the distributed engine."),
    ], [1.25, 2.55, 2.7], 8)
    heading(doc, "Architecture to draw on the board")
    code(doc, "Google COVID CSVs\n       | download\n       v\nLocal dataset/  --upload-->  HDFS /covid/input\n                                 |\n                  +--------------+--------------+\n                  |                             |\n            DataNode worker 1             DataNode worker 2\n                  |                             |\n                  +--------- Spark job --------+\n                                 |\n                                 v\n                     HDFS /covid/output (Parquet)\n                                 |\n                     local output/ or MongoDB\n                                 |\n                                 v\n                         Streamlit dashboard")
    heading(doc, "How HDFS proves distribution")
    para(doc, "The NameNode manages file-system metadata. The two DataNodes store block replicas. The epidemiology.csv file is about 521 MB and was verified as four HDFS blocks at replication factor two. fsck listed both worker DataNodes for each of those blocks. This is evidence of HDFS distribution; simply mounting the same CSV into each container would not be.")
    heading(doc, "How the Spark processing is designed")
    bullet(doc, "The job reads epidemiology, index, vaccination, hospitalization, and demographic data from HDFS.")
    bullet(doc, "Spark transformations prepare global totals, country summaries, daily global trends, country-date series, vaccination and hospitalization summaries, and regional aggregates.")
    bullet(doc, "The job is configured for Spark standalone mode with one executor on each Spark worker in the demonstrated run.")
    bullet(doc, "The intended output is written to HDFS /covid/output. After a successful submit, the runner synchronizes Parquet results under output/ for Streamlit fallback.")
    callout(doc, "State the result honestly", "The latest run reached the aggregation stage but lost an executor during shuffle and stalled while writing. Explain the architecture and observed HDFS evidence, but say the end-to-end Spark result was not verified in that run.")
    heading(doc, "Dashboard highlights")
    bullet(doc, "Country/region coverage is dynamic: the prior dashboard verification displayed 232 processed country records, while the source index contained 246 country-level keys.")
    bullet(doc, "The data explorer reports records, missing values, duplicates, source columns, and metric definitions. A prior view showed 232 rows, 17 missing cells, and no duplicates; these figures may be from older results.")
    bullet(doc, "The dashboard includes country selection, global filters, historical daily/weekly/monthly grouping, map metric controls, custom chart options, exports, and a persistent light/dark theme.")
    bullet(doc, "A legacy field called total_vaccinated has an unclear meaning. The dashboard does not silently call it doses or unique people; the new pipeline schema separates those definitions.")
    heading(doc, "Suggested four-minute presentation")
    for sentence in [
        "Introduction: Our project applies a Big Data pipeline to public COVID-19 records and keeps storage, distributed processing, and visualization as separate layers.",
        "Architecture: Docker Compose starts a Hadoop master and two worker nodes, plus Spark master/worker services and a Streamlit dashboard.",
        "Storage: We place the CSVs in HDFS. The NameNode tracks metadata, while the DataNodes hold replicas. The epidemiology file has four blocks replicated to both workers in the verified input check.",
        "Processing: PySpark reads HDFS data and groups it into global, country, and date-level analytics. Spark workers register executors so tasks can run across worker containers.",
        "Visualization: Streamlit reads processed summaries and lets the user compare countries, dates, metrics, and charts. It is not the processing engine.",
        "Validation: The cluster and HDFS input passed the latest check. The latest Spark job did not finish because an executor was lost during shuffle, so I will not present the existing output as fresh successful analytics.",
    ]:
        number(doc, sentence)
    heading(doc, "Useful demo commands")
    code(doc, "docker compose ps\ndocker exec hadoop-master hdfs dfsadmin -report\ndocker exec hadoop-master hdfs dfs -ls /covid/input\ndocker exec hadoop-master hdfs fsck /covid/input/epidemiology.csv -files -blocks -locations\nrun-analysis.bat -Mode cluster\ndocker exec hadoop-master hdfs dfs -ls -R /covid/output")
    heading(doc, "Likely questions and answers")
    table(doc, ["Question", "Clear answer"], [
        ("Why use HDFS?", "It stores large inputs in blocks and replicates blocks across DataNodes, providing distributed storage rather than copies mounted into containers."),
        ("Why Spark?", "The project uses PySpark DataFrame transformations and aggregations over the HDFS source tables; Streamlit is only the presentation layer."),
        ("Is this a three-machine cluster?", "No. It is a 3-node Hadoop cluster represented by containers on one Docker host. The worker roles and distributed block locations are real within that cluster."),
        ("What does replication factor two mean?", "HDFS keeps two copies of each block. The checked epidemiology blocks were located on both worker DataNodes."),
        ("Does total_vaccinated mean doses?", "Not necessarily. The legacy field has no guaranteed definition. The updated data contract separates people vaccinated from doses administered."),
        ("Did the Spark job finish?", "The latest attempt registered two worker executors and prepared aggregations but lost an executor during shuffle; it did not produce a verified successful run."),
        ("Why keep Streamlit?", "It makes processed analytics explorable through filters, charts, map controls, and data-quality views without moving distributed processing into the UI."),
    ], [2.1, 4.4], 8)
    heading(doc, "Presentation discipline")
    bullet(doc, "Show live command output. Do not manufacture terminal output or screenshots.")
    bullet(doc, "Distinguish a healthy HDFS input file from the overall report's four under-replicated block counter.")
    bullet(doc, "Call the dashboard country count a processed-result count (232 in the prior check), not the source index count (246).")
    bullet(doc, "State the Spark failure and avoid describing old output as a result of the latest run.")
    return doc


def build_report():
    doc = setup_doc("Project report")
    cover(doc, "COVID 19 Big Data Analytics", "A distributed storage and analytics platform using Hadoop, Apache Spark, and Streamlit", "Academic project report", True)
    heading(doc, "Abstract")
    para(doc, "This report presents a Docker-based COVID-19 analytics project that combines HDFS, a three-node Hadoop cluster, Apache Spark, and a Streamlit dashboard. Public CSV datasets are uploaded to HDFS, replicated across two DataNodes, and read by a PySpark job designed to produce global, country, date, vaccination, hospitalization, and regional summaries. The dashboard presents processed results and data-quality information. The infrastructure and HDFS input were verified on 2 October 2026. The latest Spark standalone attempt registered executors on both workers and prepared the aggregate computations, but an executor was lost during shuffle and the run did not finish. Accordingly, this report distinguishes verified infrastructure from unverified current analytics output.")
    p = doc.add_paragraph()
    set_run_font(p.add_run("Keywords: "), size=9, color=BLUE, bold=True)
    set_run_font(p.add_run("Big Data, HDFS, Hadoop, Apache Spark, PySpark, Docker Compose, Streamlit, COVID-19"), size=9, color=NAVY)
    heading(doc, "1. Introduction")
    para(doc, "The project demonstrates how a large public health dataset can move through distributed storage, distributed processing, and a visualization layer. The existing Python and Streamlit application is retained. Hadoop and Spark services are added as containers so the system can be inspected and demonstrated using standard cluster and HDFS commands.")
    heading(doc, "1.1 Problem statement", 2)
    para(doc, "A local-only script can calculate summaries but does not demonstrate distributed storage or processing. This project addresses that gap by placing source files in HDFS, providing two DataNode workers, and submitting PySpark computations through a Spark master with two worker containers.")
    heading(doc, "1.2 Objectives", 2)
    for item in [
        "Retain and run the existing Streamlit dashboard in Docker.",
        "Configure one Hadoop master and two workers with NameNode, ResourceManager, DataNode, and NodeManager roles.",
        "Store source CSV files in HDFS with replication across the DataNodes.",
        "Use Apache Spark to prepare COVID-19 global, country, and date-level analytics.",
        "Provide processed summaries to Streamlit without using Streamlit as the distributed engine.",
        "Document repeatable setup, verification, presentation, and operational commands.",
    ]:
        bullet(doc, item)
    heading(doc, "2. System architecture")
    code(doc, "COVID-19 Open Data CSVs\n        | download and upload\n        v\nHDFS /covid/input  (replication factor 2)\n        |\n        +--- Hadoop master: NameNode + ResourceManager\n        +--- Worker 1: DataNode + NodeManager\n        +--- Worker 2: DataNode + NodeManager\n        |\n        v\nApache Spark job on Spark master + two Spark workers\n        |\n        v\nHDFS /covid/output -> local output/ Parquet or optional MongoDB\n        |\n        v\nStreamlit dashboard")
    table(doc, ["Layer", "Implementation", "Purpose"], [
        ("Orchestration", "Docker Compose", "Starts the Hadoop, Spark, and dashboard containers on one host."),
        ("Distributed storage", "HDFS, replication factor 2", "Stores CSV input as HDFS blocks with replicas on two DataNodes."),
        ("Resource management", "YARN ResourceManager and two NodeManagers", "Provides Hadoop scheduling services; current job is submitted to Spark standalone."),
        ("Distributed processing", "Apache Spark / PySpark", "Reads HDFS tables and prepares grouped analytics."),
        ("Results", "HDFS Parquet, local output/ Parquet, optional MongoDB", "Makes processed summaries available to dashboard and inspection commands."),
        ("Visualization", "Streamlit and Plotly", "Presents KPIs, maps, charts, filters, and quality information."),
    ], [1.25, 2.25, 3.0], 8)
    para(doc, "The three Hadoop nodes are containers on one Docker host; they are not three physical hosts. HDFS block replicas and Spark executors are distributed across worker containers.")
    heading(doc, "3. Dataset and ingestion")
    para(doc, "The source is Google COVID-19 Open Data. The ingestion scripts download epidemiology.csv, vaccinations.csv, hospitalizations.csv, demographics.csv, and index.csv into the local dataset/ directory. The upload process sends these files to HDFS under /covid/input. The epidemiology CSV was measured at 520,931,512 bytes during the HDFS check.")
    heading(doc, "3.1 HDFS verification", 2)
    para(doc, "The live HDFS report showed two DataNodes. The input directory contained all five expected files. fsck reported the epidemiology file as healthy with four blocks, replication factor two, and both worker DataNodes listed for each block. The aggregate HDFS report also showed four under-replicated blocks overall, so the per-file input check and overall filesystem counters should be reported separately.")
    heading(doc, "4. Distributed processing design")
    para(doc, "The Spark job in spark/jobs/analytics_job.py reads the HDFS tables, applies cleaning and schema transformations, joins reference and metric tables, and prepares aggregations in src/analytics/covid_analyzer.py. Outputs are designed for global metrics, country summaries, daily global summaries, country-date series, vaccination, hospitalization, and regional views. The job writes Parquet results under HDFS /covid/output. Following successful completion, runner scripts copy selected result directories to output/ for the dashboard fallback. MongoDB can be configured as an optional results store.")
    heading(doc, "4.1 Metric interpretation", 2)
    para(doc, "Country and date aggregations are derived from the data rather than hard-coded KPI values. Vaccination counts require field-level interpretation: a field that measures individual doses must be labeled as doses, while unique people require an explicit people-vaccinated source field. A legacy total_vaccinated field has no guaranteed meaning in this project's old output; the dashboard labels it as undefined rather than assuming people or doses.")
    heading(doc, "5. Streamlit dashboard")
    para(doc, "The Streamlit application consumes processed MongoDB collections when available and otherwise reads local Parquet output. The interface includes dynamic KPIs, country and date filters, an interactive map view, historical grouping, country comparisons, a custom chart builder, regional views when regional summaries are available, a data explorer, CSV/Excel export, and a persistent light/dark theme. These controls transform or filter already-processed summaries; Spark remains the distributed processing layer.")
    heading(doc, "5.1 Coverage and quality snapshot", 2)
    para(doc, "A prior dashboard check displayed 232 processed country records. The source index contained 246 country-level keys (aggregation_level=0). These counts refer to different levels: the processed dashboard count depends on the available summary output and is not hard-coded. The prior dashboard quality view showed 232 records, 17 missing cells, and zero duplicate rows. These are dated observations from the available processed output, not a guarantee about a future rerun.")
    heading(doc, "6. Verification and results")
    table(doc, ["Check", "Observed result on 2 Oct 2026", "Assessment"], [
        ("Compose configuration", "docker compose config --quiet completed successfully.", "PASS"),
        ("Image build and stack start", "docker compose build and docker compose up -d succeeded; ten services were listed.", "PASS"),
        ("Hadoop node roles", "NameNode, ResourceManager, two DataNodes, and two NodeManagers were running; master services and NodeManagers reported healthy.", "PASS"),
        ("HDFS live workers", "dfsadmin -report showed two live DataNodes.", "PASS"),
        ("HDFS input replication", "Five CSVs listed; epidemiology file had four healthy blocks, replication 2, both workers for each block.", "PASS for checked input file"),
        ("Overall HDFS replication", "dfsadmin report listed four under-replicated blocks overall.", "Needs follow-up"),
        ("Dashboard service", "Container healthy; Streamlit health endpoint returned ok.", "PASS"),
        ("Spark job", "Two executors registered and aggregations prepared; executor lost during shuffle, no successful process exit verified.", "FAIL / incomplete"),
        ("Fresh analytics outputs", "Post-failure HDFS listing could not be confirmed; do not treat any existing output as fresh.", "Not verified"),
    ], [1.35, 4.0, 1.15], 7.6)
    heading(doc, "7. Limitations and next engineering work")
    for item in [
        "The latest distributed Spark analytics run did not complete; worker loss during shuffle prevented verification of fresh output.",
        "YARN cluster-mode execution is constrained because the current Hadoop NodeManager images do not include Python 3. The documented job path uses Spark standalone mode.",
        "The aggregate HDFS report listed four under-replicated blocks even though the inspected epidemiology input file passed fsck. The remaining HDFS paths should be audited after the Spark application is stopped.",
        "A full clean-volume reset was not performed during the audit because it would delete existing HDFS data and NameNode metadata. Compose build and start were tested with existing named volumes.",
        "Country map matching is distinct from source coverage. Aliases help the visualization, while source keys remain the analytical identifiers.",
        "Regional geographic boundary files are not present for every state or province, so a universal regional choropleth is unavailable.",
    ]:
        bullet(doc, item)
    heading(doc, "8. Conclusion")
    para(doc, "The project contains a genuine containerized Hadoop cluster with HDFS storage replicated across two DataNodes, Spark master/worker services, and a Streamlit visualization layer. Live verification confirmed the three-node Hadoop topology, five HDFS inputs, two live DataNodes, and healthy replicated blocks for the epidemiology file. The dashboard service also ran successfully. However, the latest Spark analytics submission lost an executor during shuffle and did not finish; therefore, the end-to-end fresh analytics result remains unverified. A successful rerun with current Parquet output is required before claiming full pipeline completion.")
    heading(doc, "References")
    refs = [
        "Google Cloud Platform. COVID-19 Open Data. https://github.com/GoogleCloudPlatform/covid-19-open-data",
        "Apache Hadoop. HDFS Architecture Guide. https://hadoop.apache.org/docs/r3.2.1/hadoop-project-dist/hadoop-hdfs/HdfsDesign.html",
        "Apache Spark. Spark 3.3.0 Documentation. https://spark.apache.org/docs/3.3.0/",
        "Streamlit. Documentation. https://docs.streamlit.io/",
        "Docker. Docker Compose Documentation. https://docs.docker.com/compose/",
    ]
    for i, ref in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f"[{i}] {ref}")
    return doc


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    if not LOGO.exists():
        raise FileNotFoundError(f"Supplied university logo was not found: {LOGO}")
    docs = [
        ("01-Run-the-Project-Step-by-Step.docx", build_run_guide()),
        ("02-Explain-the-Project-to-Your-Professor.docx", build_explainer()),
        ("03-University-Project-Report-4CS101ME25.docx", build_report()),
    ]
    for filename, document in docs:
        target = OUT / filename
        document.save(target)
        print(f"CREATED {target}")


if __name__ == "__main__":
    main()
